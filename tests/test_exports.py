import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import unittest
import pandas as pd
from openpyxl import load_workbook
from beams_ec2.exporting import write_excel, export_snapshot, start_export
from beams_ec2.release import BeamsEC2App
from tests.test_reuse import fixture, V


def export_app():
    app, results = fixture()
    app.df_results = results
    app.df_summary = results
    app.df_env = results
    app.df_calc_input = results
    for attr in ('df_failures','df_ok','df_validation','df_notes'):
        setattr(app, attr, pd.DataFrame())
    app.input_file_path = 'test'
    app._metadata_df = lambda: pd.DataFrame([['Programa','BeamsEC2']], columns=['Campo','Valor'])
    app._parameters_df = lambda: pd.DataFrame([['Test','value']], columns=['Parâmetro','Valor'])
    app.build_shortlists_df = lambda: pd.DataFrame()
    app._memory_for_report = BeamsEC2App._memory_for_report
    return app


class ExportTests(unittest.TestCase):
    def test_single_pass_keeps_every_row_and_column_and_literal_strings(self):
        app = export_app()
        app.df_clean = pd.DataFrame({'first':['=1+1','https://example.org',None], 'middle':[1.25,-2.5,3], 'last':['á','fim','última']})
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'export.xlsx';write_excel(app,path)
            wb=load_workbook(path,read_only=False,data_only=False)
            ws=wb['03_Entrada_Dados']
            self.assertEqual(list(ws.values),[('first','middle','last'),('=1+1',1.25,'á'),('https://example.org',-2.5,'fim'),(None,3,'última')])
            self.assertEqual(ws['A2'].data_type,'s')
            self.assertIsNone(ws['A3'].hyperlink)
            self.assertEqual(ws.freeze_panes,'A2')
            self.assertEqual(len(wb.sheetnames),26)
            self.assertIn('Governantes tensões aço / betão', [row[5] for row in wb['08_Memoria_Calculo'].values][1:])
            wb.close()

    def test_failure_preserves_previous_export(self):
        app=export_app()
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'export.xlsx';path.write_bytes(b'previous export')
            def broken(_):
                yield 'first', pd.DataFrame({'a':[1]})
                raise RuntimeError('source unavailable')
            with patch('beams_ec2.exporting.workbook_frames',broken):
                with self.assertRaisesRegex(RuntimeError,'source unavailable'):write_excel(app,path)
            self.assertEqual(path.read_bytes(),b'previous export')
            self.assertEqual(list(Path(folder).iterdir()),[path])

    def test_background_export_uses_frozen_data_and_keeps_ui_available(self):
        from threading import Event, get_ident
        app=export_app();ui_thread=get_ident();original=app.var_cover
        class MainThreadValue:
            def get(self):
                if get_ident()!=ui_thread:raise AssertionError('Tk read from worker')
                return original.get()
        app.var_cover=MainThreadValue()
        app.status_var=V('');app.status_var.set=lambda v:setattr(app.status_var,'x',v)
        callbacks=[];app.after=lambda ms,callback:callbacks.append(callback)
        started=Event();release=Event();done=Event();captured=[]
        def writer(snapshot,path):
            started.set();release.wait(5)
            captured.append((snapshot.var_cover.get(),snapshot.df_results.member.iloc[0]))
            done.set()
        with tempfile.TemporaryDirectory() as folder,patch.object(BeamsEC2App,'_write_excel',writer):
            start_export(app,str(Path(folder)/'test.xlsx'),'xlsx')
            self.assertTrue(started.wait(5));self.assertTrue(app._export_running)
            app.df_results.loc[0,'member']='changed';original.x=100
            callbacks.pop(0)()  # UI poll returns while the writer is still busy.
            self.assertTrue(callbacks)
            release.set();self.assertTrue(done.wait(5))
            import time
            deadline=time.monotonic()+5
            while app._export_running and time.monotonic()<deadline:
                callbacks.pop(0)();time.sleep(.001)
            self.assertFalse(app._export_running)
            self.assertEqual(captured,[(35,'1')])
            self.assertIn('XLSX exportado',app.status_var.get())
