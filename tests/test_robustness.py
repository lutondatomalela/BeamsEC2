import math
import unittest
import pandas as pd
from beams_ec2.release import BeamDesigner, _prepare_input_frames
from beams_ec2.base import BeamSection
from beams_ec2.geometry import SectionProfileGeometry


def sample(**changes):
    row = {'Member/Node/Case': '1/1/101 (C)', 'Station (m)': 0,
           'MY (kNm)': 60, 'MZ (kNm)': 0, 'MX (kNm)': 0,
           'FZ (kN)': 40, 'FX (kN)': 0, 'Length (m)': 6,
           'HY (cm)': 30, 'HZ (cm)': 60, 'Material': 'C30/37'}
    row.update(changes)
    return pd.DataFrame([row])


class InputTests(unittest.TestCase):
    def test_nonfinite_actions(self):
        for value in ['abc', float('nan'), float('inf'), '-inf']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                _prepare_input_frames(sample(**{'MY (kNm)': value}))

    def test_missing_action_column(self):
        with self.assertRaises(ValueError):
            _prepare_input_frames(sample().drop(columns='FZ (kN)'))

    def test_geometry_and_station(self):
        for change in [{'HY (cm)': 0}, {'Length (m)': -1}, {'Station (m)': 7}]:
            with self.subTest(change=change), self.assertRaises(ValueError):
                _prepare_input_frames(sample(**change))

    def test_variable_section_rejected(self):
        df = pd.concat([sample(), sample(**{'HZ (cm)': 70})], ignore_index=True)
        with self.assertRaises(ValueError):
            _prepare_input_frames(df)

    def test_invalid_material(self):
        for value in ['S355', 'C30/99', 'C100/115']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                _prepare_input_frames(sample(**{'Material': value}))

    def test_duplicate_dataframe_index(self):
        df = pd.concat([sample(), sample(**{'Station (m)': 3, 'MY (kNm)': 100})])
        _, env = _prepare_input_frames(df)
        self.assertEqual(env.iloc[0].m_pos_ed_kNm, 100)

    def test_axes(self):
        _, env = _prepare_input_frames(sample(), moment_axis='mz', shear_axis='fz')
        self.assertEqual(env.iloc[0].biaxial_m_ed_kNm, 60)


class CalculationTests(unittest.TestCase):
    def setUp(self):
        self.d = BeamDesigner()
        self.section = BeamSection(300, 600, 0, 0)
        self.fyd = 500 / 1.15

    def test_singly_reinforced_closed_form(self):
        As, depth, fcd = 1200, 550, 20
        a = As * self.fyd / (300 * fcd)
        expected = As * self.fyd * (depth - a / 2) / 1e6
        result = self.d.flexural_capacity(As, 0, self.section, depth, 50, fcd, self.fyd)
        self.assertAlmostEqual(result['MRd_kNm'], expected, places=8)
        self.assertLess(abs(result['equilibrium_residual_N']), 1e-6)

    def test_opposite_face_steel_not_assumed_yielding(self):
        r = self.d.flexural_capacity(1000, 3000, self.section, 550, 50, 20, self.fyd)
        self.assertLess(abs(r['sigma_c_MPa']), self.fyd)
        self.assertLess(abs(r['equilibrium_residual_N']), 1e-6)
        self.assertGreater(r['MRd_kNm'], 0)

    def test_overreinforced_tension_not_yielding(self):
        r = self.d.flexural_capacity(15000, 0, self.section, 550, 50, 20, self.fyd)
        self.assertLess(r['sigma_t_MPa'], self.fyd)

    def test_profile_and_rectangle_agree(self):
        p = SectionProfileGeometry(600, 300, 300, 0, 300, 0)
        args = (1200, 500)
        r = self.d.flexural_capacity(*args, self.section, 550, 50, 20, self.fyd)
        g = self.d.flexural_capacity(*args, p, 550, 50, 20, self.fyd)
        self.assertAlmostEqual(r['MRd_kNm'], g['MRd_kNm'], places=10)

    def test_torsion_uses_two_leg_area(self):
        r = self.d.torsion_requirements(20, self.section, 30, 20, self.fyd)
        expected = 20e6 / (r['Ak_mm2'] * self.fyd * self.d.cot_theta)
        self.assertAlmostEqual(r['Asw_s_torsion_req_mm2_per_mm'], expected)
        self.assertEqual(r['cot_theta_torsion'], self.d.cot_theta)

    def test_small_equilibrium_torsion_kept(self):
        r = self.d.torsion_requirements(0.01, self.section, 30, 20, self.fyd)
        self.assertEqual(r['torsion_considered'], 'Sim')
        self.assertGreater(r['Asl_torsion_req_mm2'], 0)

    def test_invalid_constructor(self):
        for kw in [{'gamma_c': 0}, {'cover_mm': -1}, {'cot_theta': 3}, {'fyk': math.inf}]:
            with self.subTest(kw=kw), self.assertRaises(ValueError):
                BeamDesigner(**kw)

    def design(self, **change):
        _, env = _prepare_input_frames(sample(**change))
        return self.d.design_one(env.iloc[0])

    def test_axial_requires_review(self):
        r = self.design(**{'FX (kN)': 100})
        self.assertNotEqual(r['status'], 'OK')
        self.assertIn('N-M', r['recommendations'])

    def test_biaxial_requires_review(self):
        r = self.design(**{'MZ (kNm)': 10})
        self.assertNotEqual(r['status'], 'OK')
        self.assertIn('biaxial', r['recommendations'])

    def test_high_strength_requires_review(self):
        r = self.design(**{'Material': 'C60/75'})
        self.assertNotEqual(r['status'], 'OK')
        self.assertIn('alta resistência', r['recommendations'])

    def test_combined_strut_failure(self):
        initial = self.design(**{'MX (kNm)': 1})
        r = self.design(**{'FZ (kN)': 0.7 * initial['VRd_max_kN'],
                           'MX (kNm)': 0.7 * initial['TRd_max_kNm']})
        self.assertGreater(r['eta_vt'], 1)
        self.assertEqual(r['status'], 'Falha')
        self.assertIn('6.29', r['failure_reason'])

    def test_closed_stirrups_only_two_legs(self):
        r = self.d.choose_stirrups(0.5, 300, 600, 550, torsion=True)
        self.assertEqual(r['stirrup_legs'], 2)

    def test_invalid_envelope_not_zero(self):
        _, env = _prepare_input_frames(sample())
        env.loc[0, 'v_ed_kN'] = float('nan')
        with self.assertRaises(ValueError):
            self.d.design_one(env.iloc[0])


class ServiceabilityTests(unittest.TestCase):
    def test_partial_diagram_not_stretched(self):
        from beams_ec2.serviceability import _moment_diagram_serviceability
        group = pd.DataFrame({'station': [0, 0.5, 1], 'my': [0, 10, 0]})
        x, _, quality, _ = _moment_diagram_serviceability(group, 6, 'my')
        self.assertEqual(quality, 'Insuficiente')
        self.assertEqual(len(x), 0)

    def test_full_diagram_preserves_units(self):
        from beams_ec2.serviceability import _moment_diagram_serviceability
        group = pd.DataFrame({'station': [6, 0, 3], 'my': [0, 0, 10]})
        x, M, quality, _ = _moment_diagram_serviceability(group, 6, 'my')
        self.assertEqual(list(x), [0, 3, 6])
        self.assertEqual(list(M), [0, 10, 0])
        self.assertEqual(quality, 'Aceitável')

    def test_conflicting_station_not_averaged(self):
        from beams_ec2.serviceability import _moment_diagram_serviceability
        group = pd.DataFrame({'station': [0, 3, 3, 6], 'my': [0, 10, 50, 0]})
        x, _, quality, _ = _moment_diagram_serviceability(group, 6, 'my')
        self.assertEqual(len(x), 0)
        self.assertEqual(quality, 'Insuficiente')

    def test_local_rebar_positions_not_rescaled(self):
        from beams_ec2.advanced import _group_positions_advanced
        group = pd.DataFrame({'station': [0.5, 0, 1]})
        _, x = _group_positions_advanced(group, 6)
        self.assertEqual(list(x), [0, 0.5, 1])


class ReportingTests(unittest.TestCase):
    def test_selftest_on_cp1252_console(self):
        import os
        import subprocess
        import sys
        from pathlib import Path
        result = subprocess.run([sys.executable, "selftest.py"],
                                cwd=Path(__file__).resolve().parents[1],
                                env={**os.environ, "PYTHONIOENCODING": "cp1252:strict"},
                                capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr.decode("ascii", errors="replace"))

    def test_summary_does_not_hide_review_or_mix_cases(self):
        from beams_ec2.base import build_summary_by_member
        rows = pd.DataFrame([
            dict(member='1', story='', name='', case='101', status='OK', eta_global=0.9, detail=99),
            dict(member='1', story='', name='', case='102', status='Verificar', eta_global=0.2, detail=None)])
        summary = build_summary_by_member(rows)
        self.assertEqual(summary.iloc[0]['case'], '102')
        self.assertTrue(pd.isna(summary.iloc[0]['detail']))

    def test_no_scalar_case_reduction(self):
        from beams_ec2.base import reduce_to_governing_cases
        data = pd.DataFrame({'case': ['101', '102', '103']})
        self.assertEqual(len(reduce_to_governing_cases(data)), 3)

    def test_xlsx_roundtrip_all_audit_sheets(self):
        from types import SimpleNamespace
        from tempfile import TemporaryDirectory
        from pathlib import Path
        from openpyxl import load_workbook
        from beams_ec2.release import BeamsEC2App
        clean, env = _prepare_input_frames(sample(**{'MX (kNm)': 1}))
        results = BeamDesigner().design_dataframe(env)
        app = SimpleNamespace(df_results=results, df_env=env, df_clean=clean,
                              df_calc_input=env, df_summary=results,
                              df_failures=results[results.status == 'Falha'],
                              df_ok=results[results.status == 'OK'],
                              df_validation=pd.DataFrame(), df_notes=pd.DataFrame(),
                              _metadata_df=lambda: pd.DataFrame({'Campo':['Programa'], 'Valor':['BeamsEC2']}),
                              _parameters_df=lambda: pd.DataFrame(),
                              build_shortlists_df=lambda: pd.DataFrame())
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'audit.xlsx'
            BeamsEC2App._write_excel(app, str(path))
            book = load_workbook(path, read_only=True)
            expected = {'05A_Geometria', '06_Resultados_Completos', '10_Corte_Torcao',
                        '11A_ELS_Pontos', '11C_ELS_Combinacoes', '12A_Porm_Longitudinal'}
            self.assertTrue(expected.issubset(book.sheetnames))
            header = [c.value for c in next(book['06_Resultados_Completos'].rows)]
            self.assertIn('eta_vt', header)
            self.assertEqual(book['06_Resultados_Completos'].max_row, 2)
            book.close()


if __name__ == '__main__':
    unittest.main()
